
"""
NAWI Test Report DOCX generator.

Consumes the same structured report-data dictionary used by the NAWI PDF
generator and produces a print-ready Word (.docx) report with the same
visual structure, section order, data fields, status treatment, and
report conventions.

The generator is a presentation layer only. It does not calculate OIML
measurement results.
"""

import io
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from backend.app.services.verification.qr_utils import generate_qr_png_bytes


LOGO_PATH = Path(__file__).resolve().parent / "assets" / "logo.png"

# ============================================================
# THEME - matches the finalized PDF generator
# ============================================================

BURGUNDY = "6B2D2D"
BURGUNDY_DARK = "4B2020"
BURGUNDY_LIGHT = "F4EAEA"
CHARCOAL = "2E2E2E"
TEXT = "252525"
MUTED = "666666"
BORDER = "B7B0B0"
ROW_ALT = "FAF7F7"
WHITE = "FFFFFF"
GREEN = "2F6B3A"
RED = "9B3030"
GREY = "6E6E6E"

FONT_NAME = "Calibri"
BODY_SIZE = Pt(9.5)

PAGE_WIDTH = Inches(8.27)
PAGE_HEIGHT = Inches(11.69)
PAGE_MARGIN = Inches(0.78)
HEADER_DISTANCE = Inches(0.34)
FOOTER_DISTANCE = Inches(0.34)

USABLE_WIDTH = PAGE_WIDTH - (PAGE_MARGIN * 2)

IST = timezone(timedelta(hours=5, minutes=30), name="IST")
UTC = timezone.utc


# ============================================================
# DATA HELPERS
# ============================================================

def safe(value: Any, default: str = "Not recorded") -> str:
    if value is None:
        return default
    if isinstance(value, bool):
        return "Yes" if value else "No"
    text = str(value).strip()
    return text if text else default


def optional(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def first_present(*values: Any, default: Any = None) -> Any:
    for value in values:
        if value is not None and value != "":
            return value
    return default


def format_value(value: Any) -> str:
    if value is None:
        return "Not recorded"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:g}"
    return str(value)


def parse_datetime(value: Any) -> datetime | None:
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        dt = value
    elif isinstance(value, date):
        dt = datetime(value.year, value.month, value.day)
    elif isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        if text.endswith("Z"):
            text = text[:-1] + "+00:00"
        try:
            dt = datetime.fromisoformat(text)
        except ValueError:
            for fmt in (
                "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M",
                "%d-%m-%Y %H:%M:%S",
                "%d-%m-%Y %H:%M",
            ):
                try:
                    dt = datetime.strptime(text, fmt)
                    break
                except ValueError:
                    continue
            else:
                return None
    else:
        return None

    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=UTC)

    return dt.astimezone(IST)


def format_datetime(value: Any, include_time: bool = True) -> str:
    dt = parse_datetime(value)
    if dt is None:
        return "Not recorded"
    return dt.strftime("%d-%m-%Y %H:%M IST" if include_time else "%d-%m-%Y")


def display_timestamp(value: Any) -> str:
    return format_datetime(value, include_time=True)


def get_tester_name(tester: dict) -> str:
    name = first_present(tester.get("name"), tester.get("full_name"))
    if name:
        return str(name)
    full = " ".join(
        part for part in
        (tester.get("first_name") or "", tester.get("last_name") or "")
        if part
    ).strip()
    return full or "Not recorded"


def get_reviewer_name(report: dict) -> str:
    reviewer = report.get("reviewer") or {}
    name = first_present(
        reviewer.get("name"),
        report.get("reviewer_name"),
        report.get("approved_by_name"),
    )
    if name:
        return str(name)
    full = " ".join(
        part for part in
        (reviewer.get("first_name") or "", reviewer.get("last_name") or "")
        if part
    ).strip()
    return full or "Not recorded"


def get_test_result(test: dict) -> str:
    detailed = test.get("results") or []
    detailed_result = (
        detailed[0] if detailed and isinstance(detailed[0], dict) else {}
    )
    value = first_present(
        test.get("result"),
        test.get("pass_fail"),
        detailed_result.get("pass_fail"),
        default="Not recorded",
    )
    return str(value)


def get_test_applicability(test: dict) -> str:
    return str(first_present(
        test.get("applicability_status"),
        test.get("applicability"),
        default="Not recorded",
    ))


# ============================================================
# OOXML / FORMATTING HELPERS
# ============================================================

def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def _set_cell_borders(cell, color: str = BORDER, size: int = 4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _set_fixed_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_column_widths(table, widths) -> None:
    table.autofit = False
    _set_fixed_layout(table)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def _set_paragraph_spacing(paragraph, before=0, after=0, line=1.0) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def _set_run(run, *, bold=False, italic=False, color=TEXT, size=BODY_SIZE) -> None:
    run.font.name = FONT_NAME
    run.font.size = size
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def _set_cell_text(
    cell,
    text: Any,
    *,
    bold=False,
    italic=False,
    color=TEXT,
    size=BODY_SIZE,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.keep_together = True
    p.alignment = align
    _set_paragraph_spacing(p, after=0, line=1.0)
    run = p.add_run(safe(text, ""))
    _set_run(run, bold=bold, italic=italic, color=color, size=size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _status_color(value: Any) -> str:
    text = str(value).strip().upper() if value is not None else ""
    if text == "PASS":
        return GREEN
    if text == "FAIL":
        return RED
    return CHARCOAL


def _style_heading(paragraph, *, size=Pt(11.5), color=BURGUNDY, before=5, after=3) -> None:
    paragraph.paragraph_format.keep_with_next = True
    _set_paragraph_spacing(paragraph, before=before, after=after, line=1.0)
    for run in paragraph.runs:
        _set_run(run, bold=True, color=color, size=size)


def _add_page_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    _set_run(run, color=MUTED, size=Pt(8))
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


# ============================================================
# TABLE BUILDERS
# ============================================================

def _add_kv_table(document, pairs, *, label_width=Inches(1.72)):
    table = document.add_table(rows=len(pairs), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    value_width = USABLE_WIDTH - label_width
    _set_column_widths(table, [label_width, value_width])

    for idx, (label, value) in enumerate(pairs):
        row = table.rows[idx]
        _set_cell_text(row.cells[0], label, bold=True, size=Pt(8.5))
        _set_cell_shading(row.cells[0], BURGUNDY_LIGHT)
        _set_cell_text(row.cells[1], value, size=Pt(8.5))
        for cell in row.cells:
            _set_cell_borders(cell)
        _prevent_row_split(row)

    return table


def _add_four_cell_table(document, rows_data, widths, *, header=True, row_alt=True):
    table = document.add_table(rows=0, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(table, widths)

    if header:
        row = table.add_row()
        headers = rows_data.pop(0)
        for i, value in enumerate(headers):
            _set_cell_text(
                row.cells[i], value, bold=True, color=WHITE,
                size=Pt(7.8), align=WD_ALIGN_PARAGRAPH.CENTER,
            )
            _set_cell_shading(row.cells[i], BURGUNDY)
            _set_cell_borders(row.cells[i])
        _set_repeat_header(row)

    for idx, values in enumerate(rows_data):
        row = table.add_row()
        for i, value in enumerate(values):
            _set_cell_text(
                row.cells[i],
                value,
                bold=(i in (0, 2)),
                size=Pt(8.2),
            )
            if i in (0, 2):
                _set_cell_shading(row.cells[i], BURGUNDY_LIGHT)
            elif row_alt and idx % 2 == 1:
                _set_cell_shading(row.cells[i], ROW_ALT)
            _set_cell_borders(row.cells[i])
        _prevent_row_split(row)

    return table


def _add_header_table(document, headers, rows_data, widths, align_cols=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(table, widths)

    for i, header in enumerate(headers):
        _set_cell_text(
            table.rows[0].cells[i],
            header,
            bold=True,
            color=WHITE,
            size=Pt(7.6),
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        _set_cell_shading(table.rows[0].cells[i], BURGUNDY)
        _set_cell_borders(table.rows[0].cells[i])
    _set_repeat_header(table.rows[0])

    for r_idx, row_values in enumerate(rows_data):
        row = table.add_row()
        for i, value in enumerate(row_values):
            align = (
                align_cols[i]
                if align_cols and i < len(align_cols)
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            color = None
            bold = False
            if isinstance(value, tuple):
                value, color, bold = value
            _set_cell_text(
                row.cells[i],
                value,
                bold=bold,
                color=color or TEXT,
                size=Pt(8.1),
                align=align,
            )
            if r_idx % 2 == 1:
                _set_cell_shading(row.cells[i], ROW_ALT)
            _set_cell_borders(row.cells[i])
        _prevent_row_split(row)

    return table


# ============================================================
# HEADER / FOOTER
# ============================================================

def _add_page_header_footer(section, data) -> None:
    lab = data.get("laboratory") or {}
    session = data.get("test_session") or {}
    report = data.get("report") or {}

    lab_name = safe(lab.get("name"), "Testing Laboratory")
    report_number = first_present(report.get("report_number"), "Not recorded")
    application_number = first_present(session.get("application_number"), "Not recorded")

    # Header
    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_paragraph_spacing(p, after=0)

    left = p.add_run("NAWI TEST REPORT")
    _set_run(left, bold=True, color=BURGUNDY, size=Pt(7.4))
    p.add_run("    ")
    centerish = p.add_run(f"{lab_name} | {safe(report_number)} | {safe(application_number)}")
    _set_run(centerish, color=MUTED, size=Pt(7.2))

    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), BORDER)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    _set_paragraph_spacing(fp, before=0, after=0)

    left = fp.add_run("Controlled report copy")
    _set_run(left, color=MUTED, size=Pt(7.4))

    tab_p = fp.paragraph_format
    tab_stops = tab_p.tab_stops
    tab_stops.add_tab_stop(int(USABLE_WIDTH / 2))
    tab_stops.add_tab_stop(int(USABLE_WIDTH))

    mid = fp.add_run("\tPage ")
    _set_run(mid, color=MUTED, size=Pt(7.4))
    _add_page_field(fp, "PAGE")

    of_run = fp.add_run(" of ")
    _set_run(of_run, color=MUTED, size=Pt(7.4))
    _add_page_field(fp, "NUMPAGES")

    right = fp.add_run("\tNAWI Test & Compliance System")
    _set_run(right, color=MUTED, size=Pt(7.4))

    p_pr = fp._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "6")
    top.set(qn("w:space"), "3")
    top.set(qn("w:color"), BORDER)
    p_bdr.append(top)
    p_pr.append(p_bdr)


# ============================================================
# SECTION BUILDERS
# ============================================================

def _build_report_header(document, data):
    lab = data.get("laboratory") or {}
    standard = data.get("standard") or {}
    report = data.get("report") or {}
    session = data.get("test_session") or {}
    instrument = data.get("instrument") or {}

    # Verification QR gets its own column in this SAME masthead table (not a
    # separate table), so it stays inside the bordered box, to the right of
    # the laboratory name/address — matching the agreed placement. This
    # function runs once per document, so it only appears on page 1.
    QR_COL_WIDTH = Inches(0.9)
    masthead = document.add_table(rows=1, cols=3)
    masthead.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(
        masthead,
        [Inches(1.35), USABLE_WIDTH - Inches(1.35) - QR_COL_WIDTH, QR_COL_WIDTH],
    )

    logo_cell, org_cell, qr_cell = masthead.rows[0].cells
    
    logo_cell.text = ""
    logo_paragraph = logo_cell.paragraphs[0]
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    logo_run = logo_paragraph.add_run()
    logo_run.add_picture(
        str(LOGO_PATH),
        width=Inches(1.0),
    )
    
    _set_cell_borders(logo_cell, size=6)
    logo_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    org_cell.text = ""
    p = org_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(safe(lab.get("name"), "TESTING LABORATORY"))
    _set_run(r, bold=True, color=CHARCOAL, size=Pt(12))

    p2 = org_cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_spacing(p2, after=0)
    r2 = p2.add_run(safe(lab.get("address")))
    _set_run(r2, color=MUTED, size=Pt(8))

    loc = " / ".join(
        str(v) for v in
        (lab.get("city"), lab.get("state"), lab.get("pincode"))
        if v
    )
    if loc:
        p3 = org_cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_paragraph_spacing(p3, after=0)
        _set_run(p3.add_run(loc), color=MUTED, size=Pt(7.7))

    contact = " | ".join(str(v) for v in (lab.get("phone"), lab.get("email")) if v)
    if contact:
        p4 = org_cell.add_paragraph()
        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_paragraph_spacing(p4, after=0)
        _set_run(p4.add_run(contact), color=MUTED, size=Pt(7.4))

    _set_cell_borders(org_cell, size=6)

    qr_bytes = generate_qr_png_bytes(str(session.get("test_session_id") or ""))
    qr_cell.text = ""
    qr_paragraph = qr_cell.paragraphs[0]
    qr_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qr_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    qr_run = qr_paragraph.add_run()
    qr_run.add_picture(io.BytesIO(qr_bytes), width=Inches(0.75))
    _set_cell_borders(qr_cell, size=6)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_spacing(p, before=5, after=0)
    _set_run(
        p.add_run("NON-AUTOMATIC WEIGHING INSTRUMENTS"),
        bold=True, color=MUTED, size=Pt(8.5),
    )

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_spacing(p, after=1)
    _set_run(
        p.add_run("NAWI TEST REPORT"),
        bold=True, color=BURGUNDY_DARK, size=Pt(18),
    )

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_spacing(p, after=6)
    _set_run(
        p.add_run("Formal test and evaluation record"),
        italic=True, color=MUTED, size=Pt(8.5),
    )

    report_date = first_present(report.get("report_date"), report.get("generated_at"))
    overall = first_present(
        report.get("overall_result"),
        session.get("overall_result"),
        data.get("overall_result"),
    )
    standard_code = first_present(
        standard.get("standard_code"),
        standard.get("title"),
        "OIML R 76",
    )
    standard_version = first_present(
        standard.get("version"),
        standard.get("edition_year"),
        "Not recorded",
    )

    report_rows = [
        ("Report No.", report.get("report_number"), "Report Date", format_datetime(report_date, False)),
        ("Report Version", first_present(report.get("report_version"), report.get("version")), "Report Status", first_present(report.get("report_status"), report.get("status"))),
        ("Application No.", session.get("application_number"), "Test Session No.", session.get("session_number")),
        ("Applicable Standard", standard_code, "Type Designation", instrument.get("type_designation")),
        ("Standard Edition / Version", standard_version, "Overall Result", overall),
    ]

    table = document.add_table(rows=0, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(table, [Inches(1.48), Inches(1.92), Inches(1.48), Inches(1.92)])

    for row_index, values in enumerate(report_rows):
        row = table.add_row()
        for i, value in enumerate(values):
            color = _status_color(value) if i == 3 and row_index == 4 else TEXT
            _set_cell_text(
                row.cells[i],
                value,
                bold=(i in (0, 2)),
                color=color,
                size=Pt(8.1),
                align=WD_ALIGN_PARAGRAPH.CENTER if i == 3 and row_index == 4 else WD_ALIGN_PARAGRAPH.LEFT,
            )
            if i in (0, 2):
                _set_cell_shading(row.cells[i], BURGUNDY_LIGHT)
            _set_cell_borders(row.cells[i])
        _prevent_row_split(row)

    return None


def _build_laboratory_section(document, data):
    lab = data.get("laboratory") or {}
    pairs = [
        ("Laboratory Code", first_present(lab.get("laboratory_code"), lab.get("code"))),
        ("Registration No.", lab.get("registration_number")),
        ("Laboratory Name", lab.get("name")),
        ("Address", lab.get("address")),
        ("City / State / PIN", " / ".join(str(v) for v in (lab.get("city"), lab.get("state"), lab.get("pincode")) if v)),
        ("Country", lab.get("country")),
        ("Phone", lab.get("phone")),
        ("Email", lab.get("email")),
    ]
    h = document.add_heading("1. Laboratory Information", level=1)
    _style_heading(h)
    _add_kv_table(document, pairs)


def _build_instrument_section(document, data):
    instrument = data.get("instrument") or {}
    h = document.add_heading("2. Instrument Under Test", level=1)
    _style_heading(h)

    headers = ["Specification", "Value", "Specification", "Value"]
    pairs = [
        ("Instrument Code", instrument.get("instrument_code"), "Manufacturer", instrument.get("manufacturer")),
        ("Model", instrument.get("model"), "Type Designation", instrument.get("type_designation") or instrument.get("type")),
        ("Serial Number", instrument.get("serial_number"), "Instrument Type", instrument.get("instrument_type")),
        ("Category", instrument.get("category"), "Accuracy Class", instrument.get("accuracy_class")),
        ("Minimum Capacity (Min)", first_present(instrument.get("minimum_capacity"), instrument.get("min_capacity")), "Maximum Capacity (Max)", first_present(instrument.get("maximum_capacity"), instrument.get("max_capacity"))),
        ("Scale Interval (d)", first_present(instrument.get("scale_interval"), instrument.get("d")), "Verification Scale Interval (e)", first_present(instrument.get("verification_scale_interval"), instrument.get("e"))),
        ("Number of Intervals (n)", first_present(instrument.get("number_of_intervals"), instrument.get("n")), "Unit", instrument.get("unit")),
        ("Indication Type", instrument.get("indication_type"), "Software / Firmware", first_present(instrument.get("software_firmware"), instrument.get("software_version"))),
    ]
    rows = [headers] + [list(map(format_value, row)) for row in pairs]
    _add_four_cell_table(
        document,
        rows,
        [Inches(1.58), Inches(1.47), Inches(1.72), Inches(1.43)],
        header=True,
        row_alt=True,
    )


def _build_session_section(document, data):
    h = document.add_heading("3. Tester / Test Session Information", level=1)
    _style_heading(h)

    session = data.get("test_session") or {}
    report = data.get("report") or {}
    tester = data.get("tester") or {}
    overall = first_present(
        report.get("overall_result"),
        session.get("overall_result"),
        data.get("overall_result"),
    )

    rows_data = [
        [
            "Session No.",
            first_present(session.get("session_number"), session.get("test_session_number")),
            "Application No.",
            first_present(session.get("application_number"), session.get("application_no")),
        ],
        [
            "Test Type",
            session.get("test_type"),
            "Session Status",
            session.get("status"),
        ],
        [
            "Tester",
            get_tester_name(tester),
            "Designation",
            tester.get("designation"),
        ],
        [
            "Test Started",
            display_timestamp(session.get("started_at") or session.get("start_date")),
            "Test Completed",
            display_timestamp(session.get("completed_at") or session.get("completion_date")),
        ],
        [
            "Overall Result",
            overall,
            "Tester ID",
            first_present(tester.get("employee_id"), tester.get("employee_code")),
        ],
    ]

    table = document.add_table(rows=0, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(
        table,
        [Inches(1.15), Inches(2.15), Inches(1.20), Inches(1.60)],
    )

    for r_idx, values in enumerate(rows_data):
        row = table.add_row()
        for i, value in enumerate(values):
            is_result = r_idx == len(rows_data) - 1 and i == 1
            _set_cell_text(
                row.cells[i],
                value,
                bold=(i in (0, 2) or is_result),
                color=_status_color(value) if is_result else TEXT,
                size=Pt(8.0),
                align=(
                    WD_ALIGN_PARAGRAPH.CENTER
                    if is_result
                    else WD_ALIGN_PARAGRAPH.LEFT
                ),
            )
            if i in (0, 2):
                _set_cell_shading(row.cells[i], BURGUNDY_LIGHT)
            _set_cell_borders(row.cells[i])
        _prevent_row_split(row)


def _build_equipment_section(document, data):
    h = document.add_heading("4. Test Equipment", level=1)
    _style_heading(h)
    equipment = data.get("test_equipment") or []

    if not equipment:
        _add_kv_table(document, [("Record", "No test equipment information was recorded in the report data.")])
        return

    headers = ["Equipment ID", "Equipment", "Model", "Serial / Identification", "Calibration"]
    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(table, [Inches(0.92), Inches(1.72), Inches(0.98), Inches(1.52), Inches(0.99)])

    for i, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[i], header, bold=True, color=WHITE, size=Pt(7.4), align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_shading(table.rows[0].cells[i], BURGUNDY)
        _set_cell_borders(table.rows[0].cells[i])
    _set_repeat_header(table.rows[0])

    for idx, item in enumerate(equipment):
        row = table.add_row()
        values = [
            item.get("equipment_code") or item.get("identification_number"),
            item.get("equipment_name"),
            item.get("model"),
            item.get("serial_number") or item.get("identification_number"),
            item.get("calibration_status"),
        ]
        for i, value in enumerate(values):
            _set_cell_text(row.cells[i], value, size=Pt(7.7), align=WD_ALIGN_PARAGRAPH.CENTER if i in (0, 4) else WD_ALIGN_PARAGRAPH.LEFT)
            if idx % 2 == 1:
                _set_cell_shading(row.cells[i], ROW_ALT)
            _set_cell_borders(row.cells[i])
        _prevent_row_split(row)

    due = []
    for item in equipment:
        code = item.get("equipment_code") or item.get("identification_number")
        due_date = item.get("calibration_due_date")
        if code and due_date:
            due.append(f"{code}: due {format_datetime(due_date, False)}")
    if due:
        p = document.add_paragraph()
        _set_paragraph_spacing(p, before=1, after=1)
        _set_run(p.add_run(" | ".join(due)), color=MUTED, size=Pt(7.5))


def _build_environment_section(document, data):
    h = document.add_heading("5. Environmental Conditions", level=1)
    _style_heading(h)

    conditions = data.get("environmental_conditions") or []
    if not conditions:
        _add_kv_table(document, [("Record", "No environmental conditions were recorded in the report data.")])
        return

    first = conditions[0]
    last = conditions[-1]

    rows = [
        ["Parameter", "Start", "End", "Recorded At"],
        ["Temperature", f"{format_value(first.get('temperature'))} °C", f"{format_value(last.get('temperature'))} °C", display_timestamp(first.get("recorded_at"))],
        ["Relative Humidity", f"{format_value(first.get('humidity'))} %", f"{format_value(last.get('humidity'))} %", display_timestamp(last.get("recorded_at"))],
        ["Atmospheric Pressure", f"{format_value(first.get('pressure'))} hPa", f"{format_value(last.get('pressure'))} hPa", "Not recorded"],
    ]
    _add_header_table(
        document,
        rows.pop(0),
        rows,
        [Inches(1.45), Inches(1.05), Inches(1.05), Inches(2.34)],
        align_cols=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )

    monitoring_source = first_present(first.get("source"), last.get("source"), default="")
    remarks = first_present(last.get("remarks"), first.get("remarks"), default="")
    parts = []
    if monitoring_source:
        parts.append(f"Monitoring source: {monitoring_source}")
    if remarks:
        parts.append(f"Remarks: {remarks}")
    if parts:
        p = document.add_paragraph()
        _set_paragraph_spacing(p, before=1, after=1)
        _set_run(p.add_run(" | ".join(parts)), color=MUTED, size=Pt(7.5))


def _build_summary_section(document, data):
    h = document.add_heading("6. Summary of Test Results", level=1)
    _style_heading(h)

    tests = data.get("tests") or []
    if not tests:
        _add_kv_table(document, [("Record", "No test results were recorded.")])
        return

    headers = ["No.", "Code", "Test / Examination", "Applicability", "Result", "Remarks"]
    table = document.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(0.4), Inches(0.86), Inches(2.12), Inches(1.05), Inches(0.7), Inches(1.11)]
    _set_column_widths(table, widths)

    for i, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[i], header, bold=True, color=WHITE, size=Pt(7.1), align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_shading(table.rows[0].cells[i], BURGUNDY)
        _set_cell_borders(table.rows[0].cells[i])
    _set_repeat_header(table.rows[0])

    counts = {"APPLICABLE": 0, "PASS": 0, "FAIL": 0, "N/A": 0}
    for idx, test in enumerate(tests, start=1):
        applicability = get_test_applicability(test)
        result = get_test_result(test)
        if applicability.upper() == "APPLICABLE":
            counts["APPLICABLE"] += 1
        if result.upper() in ("PASS", "FAIL", "N/A"):
            counts[result.upper()] += 1
        row = table.add_row()
        values = [
            idx,
            test.get("test_code"),
            test.get("test_name") or test.get("test_code"),
            applicability,
            result,
            first_present(test.get("remarks"), "Not recorded"),
        ]
        for i, value in enumerate(values):
            is_result = i == 4
            _set_cell_text(
                row.cells[i],
                value,
                bold=is_result,
                color=_status_color(value) if is_result else TEXT,
                size=Pt(7.4),
                align=WD_ALIGN_PARAGRAPH.CENTER if i in (0, 1, 3, 4) else WD_ALIGN_PARAGRAPH.LEFT,
            )
            if idx % 2 == 0:
                _set_cell_shading(row.cells[i], ROW_ALT)
            _set_cell_borders(row.cells[i])
        _prevent_row_split(row)

    p = document.add_paragraph()
    _set_paragraph_spacing(p, before=1, after=0)
    _set_run(
        p.add_run(
            f"Tests recorded: {len(tests)} | Applicable: {counts['APPLICABLE']} | "
            f"PASS: {counts['PASS']} | FAIL: {counts['FAIL']} | N/A: {counts['N/A']}"
        ),
        color=MUTED,
        size=Pt(7.7),
    )


def _build_test_metadata(document, test):
    detailed = test.get("results") or []
    detailed_result = detailed[0] if detailed and isinstance(detailed[0], dict) else {}
    name = test.get("test_name") or test.get("test_code") or "Unnamed Test"
    h = document.add_heading(
        f"Test {test.get('_index', '')} - {name} ({test.get('test_code') or 'Not recorded'})",
        level=2,
    )
    _style_heading(h, size=Pt(10.6), before=5, after=3)

    metadata = [
        ["Test Code", test.get("test_code"), "Reference Clause", test.get("reference_clause")],
        ["Applicability", get_test_applicability(test), "Category", test.get("category")],
        ["Procedure", test.get("procedure"), "Stored Result", get_test_result(test)],
    ]

    table = document.add_table(rows=0, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(table, [Inches(1.1), Inches(2.05), Inches(1.18), Inches(1.91)])

    for r_idx, values in enumerate(metadata):
        row = table.add_row()
        for i, value in enumerate(values):
            is_result = values[0] == "Applicability" and i == 3
            _set_cell_text(
                row.cells[i],
                value,
                bold=(i in (0, 2) or is_result),
                color=_status_color(value) if is_result else TEXT,
                size=Pt(7.7),
                align=WD_ALIGN_PARAGRAPH.CENTER if is_result else WD_ALIGN_PARAGRAPH.LEFT,
            )
            if i in (0, 2):
                _set_cell_shading(row.cells[i], BURGUNDY_LIGHT)
            _set_cell_borders(row.cells[i])
        _prevent_row_split(row)

    return detailed_result


def _build_observations(document, observations):
    h = document.add_heading("Observations", level=3)
    _style_heading(h, size=Pt(8.9), before=4, after=2)

    if not observations:
        p = document.add_paragraph("No observation data recorded for this test.")
        _set_paragraph_spacing(p, after=1)
        _set_run(p.runs[0], color=MUTED, size=Pt(7.8))
        return

    rows = []
    for obs in observations:
        value = first_present(
            obs.get("value_numeric"),
            obs.get("value"),
            obs.get("observed_value"),
            obs.get("value_text"),
            default="Not recorded",
        )
        rows.append([
            obs.get("parameter_name") or obs.get("parameter") or obs.get("parameter_code") or obs.get("code"),
            obs.get("parameter_code") or obs.get("code"),
            format_value(value),
            obs.get("unit"),
        ])

    _add_header_table(
        document,
        ["Parameter", "Code", "Observed Value", "Unit"],
        rows,
        [Inches(2.55), Inches(0.82), Inches(2.05), Inches(0.82)],
        align_cols=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )


def _build_calculations(document, calculations):
    h = document.add_heading("Calculations", level=3)
    _style_heading(h, size=Pt(8.9), before=4, after=2)

    if not calculations:
        p = document.add_paragraph("No calculation data recorded for this test.")
        _set_paragraph_spacing(p, after=1)
        _set_run(p.runs[0], color=MUTED, size=Pt(7.8))
        return

    rows = []
    for calc in calculations:
        value = first_present(calc.get("calculated_value"), calc.get("value"), default="Not recorded")
        rows.append([
            calc.get("calculation_type") or calc.get("calculation"),
            format_value(value),
            calc.get("unit"),
            calc.get("formula") or calc.get("calculation_type"),
        ])

    _add_header_table(
        document,
        ["Calculation", "Calculated Value", "Unit", "Formula"],
        rows,
        [Inches(1.75), Inches(1.22), Inches(0.72), Inches(2.55)],
        align_cols=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )


def _build_evaluation(document, test, detailed_result):
    h = document.add_heading("Evaluation", level=3)
    _style_heading(h, size=Pt(8.9), before=4, after=2)

    requirement = first_present(
        test.get("requirement"),
        test.get("acceptance_condition"),
        detailed_result.get("acceptance_condition"),
        default="Not recorded",
    )
    mpe = first_present(
        test.get("mpe"),
        test.get("mpe_value"),
        detailed_result.get("mpe_value"),
        default="Not recorded",
    )
    measured = detailed_result.get("measured_value")
    error = detailed_result.get("error_value")
    corrected = detailed_result.get("corrected_error")
    result = detailed_result.get("pass_fail") or get_test_result(test)

    row = [
        f"{safe(requirement)} | MPE: {format_value(mpe)}",
        format_value(measured),
        format_value(error),
        format_value(corrected),
        result,
    ]

    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(table, [Inches(2.95), Inches(1.05), Inches(0.8), Inches(1.0), Inches(0.94)])

    for i, header in enumerate(["Requirement / MPE", "Measured", "Error", "Corrected Error", "Result"]):
        _set_cell_text(table.rows[0].cells[i], header, bold=True, color=WHITE, size=Pt(7.2), align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_shading(table.rows[0].cells[i], BURGUNDY)
        _set_cell_borders(table.rows[0].cells[i])
    _set_repeat_header(table.rows[0])

    row_cells = table.add_row().cells
    for i, value in enumerate(row):
        _set_cell_text(
            row_cells[i],
            value,
            bold=(i == 4),
            color=_status_color(value) if i == 4 else TEXT,
            size=Pt(7.5),
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        _set_cell_borders(row_cells[i])
        if i == 0:
            _set_cell_shading(row_cells[i], BURGUNDY_LIGHT)
    _prevent_row_split(table.rows[1])

    acceptance = detailed_result.get("acceptance_condition") or test.get("acceptance_condition")
    summary = detailed_result.get("result_summary") or test.get("result_summary")

    extras = []
    if acceptance:
        extras.append(("Acceptance Condition", acceptance))
    if summary:
        extras.append(("Result Summary", summary))

    if extras:
        _add_kv_table(document, extras, label_width=Inches(1.48))

    calc_ver = detailed_result.get("calculation_version") or test.get("calculation_version")
    ruleset_ver = test.get("ruleset_version")
    versions = []
    if calc_ver:
        versions.append(f"Calculation version: {safe(calc_ver)}")
    if ruleset_ver:
        versions.append(f"Ruleset version: {safe(ruleset_ver)}")
    if versions:
        p = document.add_paragraph()
        _set_paragraph_spacing(p, before=1, after=1)
        _set_run(p.add_run(" | ".join(versions)), color=MUTED, size=Pt(7.4))

    remarks = test.get("remarks")
    if remarks:
        p = document.add_paragraph()
        _set_paragraph_spacing(p, before=0, after=1)
        r1 = p.add_run("Remarks: ")
        _set_run(r1, bold=True, color=TEXT, size=Pt(7.6))
        _set_run(p.add_run(str(remarks)), color=TEXT, size=Pt(7.6))


def _build_na_test(document, test, detailed_result):
    applicability = get_test_applicability(test)
    reason = first_present(test.get("na_reason"), "Not applicable")
    _add_kv_table(
        document,
        [
            ("Applicability", applicability),
            ("N/A Reason", reason),
            ("Result", "N/A"),
            ("Remarks", first_present(test.get("remarks"), "Not recorded")),
        ],
        label_width=Inches(1.48),
    )


def _build_detailed_results(document, data):
    h = document.add_heading("7. Detailed Test Results", level=1)
    _style_heading(h)

    tests = data.get("tests") or []
    for idx, raw_test in enumerate(tests, start=1):
        test = dict(raw_test)
        test["_index"] = idx
        detailed_result = _build_test_metadata(document, test)

        if str(get_test_applicability(test)).upper() not in {
            "APPLICABLE", "YES", "TRUE"
        }:
            _build_na_test(document, test, detailed_result)
        else:
            _build_observations(document, test.get("observations") or [])
            _build_calculations(document, test.get("calculations") or [])
            _build_evaluation(document, test, detailed_result)


def _build_construction_examination(document, data):
    items = data.get("construction_examination") or []
    if not items:
        return

    h = document.add_heading("Construction Examination", level=2)
    _style_heading(h, size=Pt(10.2), before=5, after=3)

    rows = []
    for item in items:
        rows.append([
            item.get("item"),
            item.get("status"),
            item.get("remarks"),
        ])

    _add_header_table(
        document,
        ["Examination Item", "Status", "Remarks"],
        rows,
        [Inches(3.0), Inches(0.9), Inches(3.35)],
        align_cols=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )


def _build_conclusion(document, data):
    h = document.add_heading("8. Overall Evaluation / Conclusion", level=1)
    _style_heading(h)

    session = data.get("test_session") or {}
    report = data.get("report") or {}
    instrument = data.get("instrument") or {}
    standard = data.get("standard") or {}

    overall = first_present(
        report.get("overall_result"),
        session.get("overall_result"),
        data.get("overall_result"),
        default="Not recorded",
    )

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(table, [Inches(2.0), USABLE_WIDTH - Inches(2.0)])
    _set_cell_text(table.rows[0].cells[0], "OVERALL RESULT", bold=True, color=CHARCOAL, size=Pt(8.5))
    _set_cell_shading(table.rows[0].cells[0], BURGUNDY_LIGHT)
    _set_cell_text(
        table.rows[0].cells[1],
        overall,
        bold=True,
        color=_status_color(overall),
        size=Pt(9),
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _set_cell_borders(table.rows[0].cells[0])
    _set_cell_borders(table.rows[0].cells[1])

    manufacturer = instrument.get("manufacturer") or "the manufacturer"
    model = instrument.get("model") or "the instrument"
    designation = instrument.get("type_designation") or ""
    standard_name = first_present(
        standard.get("standard_code"),
        standard.get("title"),
        default="the applicable standard",
    )

    narrative = (
        f"The test record presented in this report covers the recorded evaluation of "
        f"{manufacturer} {model} {designation}. The recorded results are assessed "
        f"against {standard_name}. The conclusion shown above reflects the stored "
        f"session/report result and does not independently recalculate measurement rules."
    )
    p = document.add_paragraph(narrative)
    _set_paragraph_spacing(p, before=3, after=2, line=1.05)
    _set_run(p.runs[0], color=TEXT, size=Pt(8.1))

    remarks = first_present(
        report.get("remarks"),
        session.get("remarks"),
        data.get("remarks"),
        default="",
    )
    if remarks:
        p = document.add_paragraph()
        _set_paragraph_spacing(p, after=1)
        r1 = p.add_run("Remarks: ")
        _set_run(r1, bold=True, color=TEXT, size=Pt(7.9))
        _set_run(p.add_run(str(remarks)), color=TEXT, size=Pt(7.9))


def _build_remarks(document, data):
    h = document.add_heading("9. Remarks / Non-Conformities", level=1)
    _style_heading(h)

    report = data.get("report") or {}
    session = data.get("test_session") or {}
    general = first_present(
        report.get("remarks"),
        session.get("remarks"),
        data.get("remarks"),
        default="No general remarks recorded.",
    )
    non_conformities = data.get("non_conformities") or []
    if not non_conformities:
        nc_text = "None recorded."
    else:
        descriptions = []
        for item in non_conformities:
            if isinstance(item, dict):
                descriptions.append(safe(item.get("description") or item.get("remarks"), "Not specified"))
            else:
                descriptions.append(safe(item, "Not specified"))
        nc_text = "; ".join(descriptions)

    _add_kv_table(
        document,
        [("General Remarks", general), ("Non-Conformities", nc_text)],
        label_width=Inches(1.48),
    )


def _build_approval(document, data):
    h = document.add_heading("10. Review / Approval", level=1)
    _style_heading(h)

    report = data.get("report") or {}
    reviewer = report.get("reviewer") or {}
    reviewer_name = get_reviewer_name(report)
    reviewer_designation = first_present(
        reviewer.get("designation"),
        report.get("reviewer_designation"),
        default="Not recorded",
    )
    status = str(first_present(
        report.get("report_status"),
        report.get("status"),
        default="GENERATED",
    )).upper()
    generated_on = first_present(report.get("generated_at"), report.get("report_date"))
    generated_by = first_present(report.get("generated_by"), "NAWI Test & Compliance System")
    approved_at = first_present(report.get("approved_at"), report.get("approval_date"))
    signature_status = "APPROVED" if status == "APPROVED" else "PENDING APPROVAL"

    pairs = [
        ("Report Status", status),
        ("Reviewer", reviewer_name),
        ("Reviewer Designation", reviewer_designation),
        ("Generated On", display_timestamp(generated_on)),
        ("Generated By", generated_by),
        ("Approval Date", display_timestamp(approved_at)),
        ("Signature Status", signature_status),
        ("Signed At", display_timestamp(approved_at) if approved_at else "Not recorded"),
        ("Report Hash", first_present(report.get("report_hash"), report.get("hash"))),
    ]
    _add_kv_table(document, pairs)

    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_column_widths(table, [USABLE_WIDTH / 2, USABLE_WIDTH / 2])

    headers = ["Reviewer / Approving Authority", "Signature / Seal"]
    for i, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[i], header, bold=True, size=Pt(8))
        _set_cell_shading(table.rows[0].cells[i], BURGUNDY_LIGHT)
        _set_cell_borders(table.rows[0].cells[i])

    date_text = format_datetime(approved_at, False) if approved_at else "Not recorded"
    left = table.rows[1].cells[0]
    right = table.rows[1].cells[1]
    _set_cell_text(left, f"Name: {reviewer_name}\nDate: {date_text}", size=Pt(8.1))
    _set_cell_text(right, "____________________________", size=Pt(8.1))
    left.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    right.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    for cell in (left, right):
        _set_cell_borders(cell)
        cell.height = Inches(0.65)


def _build_synthetic_notice(document, data):
    h = document.add_heading("11. Synthetic Demonstration Data Notice", level=1)
    _style_heading(h)

    notice = (
        "Where this prototype is demonstrated with synthetic or sample values, those values are for "
        "software demonstration and report-format verification only. They are not actual laboratory "
        "measurements, calibration records, or legal-metrology evidence. Production reports should be "
        "generated only from the approved structured test data recorded by the laboratory system."
    )
    table = document.add_table(rows=1, cols=1)
    _set_column_widths(table, [USABLE_WIDTH])
    cell = table.rows[0].cells[0]
    _set_cell_text(cell, notice, size=Pt(7.6))
    _set_cell_shading(cell, ROW_ALT)
    _set_cell_borders(cell)


def _build_attachments(document, data):
    h = document.add_heading("12. Attachments", level=1)
    _style_heading(h)

    attachments = data.get("attachments") or []
    rows = []
    if attachments:
        for idx, item in enumerate(attachments, start=1):
            rows.append([
                item.get("attachment_id") or idx,
                item.get("file_name") or item.get("filename"),
                item.get("attachment_type") or item.get("type"),
                item.get("description"),
            ])
    else:
        rows.append(["", "No attachments recorded.", "", ""])

    _add_header_table(
        document,
        ["ID", "File", "Type", "Description"],
        rows,
        [Inches(0.55), Inches(2.25), Inches(1.25), Inches(2.65)],
        align_cols=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )


# ============================================================
# DOCUMENT SETUP / ENTRY POINT
# ============================================================

def _apply_document_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_SIZE
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0

    for style_name, size, color in (
        ("Title", 18, BURGUNDY_DARK),
        ("Heading 1", 11.5, BURGUNDY),
        ("Heading 2", 10.6, CHARCOAL),
        ("Heading 3", 8.9, CHARCOAL),
    ):
        style = document.styles[style_name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(5)
        style.paragraph_format.space_after = Pt(3)


def create_report(data: dict, output_path: str = "NAWI_Test_Report.docx"):
    """Generate the Word report from structured report data."""
    if not isinstance(data, dict):
        raise TypeError("Report data must be a dictionary.")

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = PAGE_MARGIN
    section.bottom_margin = PAGE_MARGIN
    section.left_margin = PAGE_MARGIN
    section.right_margin = PAGE_MARGIN
    section.header_distance = HEADER_DISTANCE
    section.footer_distance = FOOTER_DISTANCE

    _apply_document_styles(document)
    _add_page_header_footer(section, data)

    # PAGE 1: IDENTIFICATION
    _build_report_header(document, data)
    _build_laboratory_section(document, data)
    _build_instrument_section(document, data)
    _build_session_section(document, data)

    # PAGE 2: EQUIPMENT + ENVIRONMENT + SUMMARY
    document.add_page_break()
    _build_equipment_section(document, data)
    _build_environment_section(document, data)
    _build_summary_section(document, data)

    # PAGE 3+: DETAILED TESTS
    document.add_page_break()
    _build_detailed_results(document, data)

    # FINAL PAGE
    document.add_page_break()
    _build_construction_examination(document, data)
    _build_conclusion(document, data)
    _build_remarks(document, data)
    _build_approval(document, data)
    _build_synthetic_notice(document, data)
    _build_attachments(document, data)

    # Embedded verification identifier, parsed back out by
    # backend/app/api/verify/routes.py when someone uploads this DOCX to
    # check it. Unlike the PDF, this is NOT cryptographically signed —
    # python-docx has no equivalent to sign_pdf_file's tamper detection —
    # so this only confirms the file corresponds to a real report, not
    # that it's unmodified. The verify page says so explicitly.
    verify_test_session_id = (data.get("test_session") or {}).get("test_session_id")
    verify_report_number = (data.get("report") or {}).get("report_number")

    document.core_properties.title = "NAWI Test Report"
    document.core_properties.subject = f"NAWI-VERIFY:{verify_test_session_id}:{verify_report_number}"
    document.core_properties.author = "NAWI Test & Compliance System"

    document.save(output)
    return str(output.resolve())


# ============================================================
# OPTIONAL LOCAL TEST
# ============================================================

if __name__ == "__main__":
    try:
        try:
            from .sample_report_data import SAMPLE_REPORT_DATA
        except ImportError:
            from sample_report_data import sample_report_data as SAMPLE_REPORT_DATA

        result = create_report(
            SAMPLE_REPORT_DATA,
            output_path="NAWI_Test_Report.docx",
        )
        print("DOCX generated successfully:")
        print(result)
    except ImportError as error:
        print("Could not import sample report data.")
        print(error)
